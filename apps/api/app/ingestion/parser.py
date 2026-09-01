import io
import logging
from dataclasses import dataclass, field

import pandas as pd

logger = logging.getLogger("vedax.parser")

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
AUDIO_EXTENSIONS = {".wav", ".mp3", ".m4a", ".ogg", ".oga", ".webm", ".flac"}
DATASET_EXTENSIONS = {".csv", ".xlsx", ".xls"}

SCANNED_PAGE_MIN_CHARS = 24
MAX_DATASET_ROWS_STORED = 100_000
MAX_RENDER_DPI = 120


@dataclass
class ParsedPage:
    number: int
    text: str
    headings: list[str] = field(default_factory=list)
    needs_ocr: bool = False
    render_png: bytes | None = None


@dataclass
class ParsedDocument:
    filename: str
    content_type: str
    source_type: str
    pages: list[ParsedPage] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)
    needs_vision: bool = False
    needs_audio: bool = False


def detect_source_type(filename: str, content_type: str) -> str:
    lower = filename.lower()
    for ext in IMAGE_EXTENSIONS:
        if lower.endswith(ext):
            return "image"
    for ext in AUDIO_EXTENSIONS:
        if lower.endswith(ext):
            return "audio"
    for ext in DATASET_EXTENSIONS:
        if lower.endswith(ext):
            return "dataset"
    for ext in PDF_EXTENSIONS | DOCX_EXTENSIONS | TEXT_EXTENSIONS:
        if lower.endswith(ext):
            return "document"
    if content_type.startswith("image/"):
        return "image"
    if content_type.startswith("audio/"):
        return "audio"
    if content_type.startswith("text/"):
        return "document"
    raise ValueError(f"unsupported file type: {filename}")


def _render_pdf_page_to_png(page, scale: float | None = None) -> bytes:
    import pypdfium2 as pdfium

    scale = scale or MAX_RENDER_DPI / 72.0
    bitmap = page.render(scale=scale)
    pil_image = bitmap.to_pil()
    buffer = io.BytesIO()
    pil_image.save(buffer, format="PNG")
    return buffer.getvalue()


def parse_pdf(data: bytes, filename: str, content_type: str) -> ParsedDocument:
    import pypdfium2 as pdfium

    parsed = ParsedDocument(filename=filename, content_type=content_type, source_type="document")
    pdf = pdfium.PdfDocument(data)
    for index in range(len(pdf)):
        page = pdf[index]
        try:
            text_page = page.get_textpage()
            text = text_page.get_text_range() or ""
        except Exception:
            text = ""
        needs_ocr = len(text.strip()) < SCANNED_PAGE_MIN_CHARS
        render_png = None
        if needs_ocr:
            try:
                render_png = _render_pdf_page_to_png(page)
            except Exception:
                logger.warning("failed rendering page %s for OCR", index + 1)
            parsed.needs_vision = True
        parsed.pages.append(
            ParsedPage(number=index + 1, text=text, needs_ocr=needs_ocr, render_png=render_png)
        )
    return parsed


def parse_docx(data: bytes, filename: str, content_type: str) -> ParsedDocument:
    import docx

    parsed = ParsedDocument(filename=filename, content_type=content_type, source_type="document")
    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    headings: list[str] = []
    for block in document.paragraphs:
        text = block.text.strip()
        if not text:
            continue
        if block.style and block.style.name and block.style.name.startswith("Heading"):
            headings.append(text)
            lines.append(f"\n## {text}\n")
        else:
            lines.append(text)
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            header = rows[0]
            records = [dict(zip(header, row)) for row in rows[1:]]
            parsed.tables.append({"columns": header, "rows": rows, "records": records[:500]})
            lines.append("\n| " + " | ".join(header) + " |\n")
            for row in rows[1:31]:
                lines.append("| " + " | ".join(row) + " |")
    parsed.pages.append(
        ParsedPage(number=1, text="\n".join(lines), headings=headings)
    )
    return parsed


def parse_text(data: bytes, filename: str, content_type: str) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    headings = [
        line.lstrip("#").strip()
        for line in text.splitlines()
        if line.startswith("#") or (line and line.isupper() and len(line) < 90)
    ]
    return ParsedDocument(
        filename=filename,
        content_type=content_type,
        source_type="document",
        pages=[ParsedPage(number=1, text=text, headings=headings)],
    )


def parse_dataset(data: bytes, filename: str, content_type: str) -> ParsedDocument:
    lower = filename.lower()
    if lower.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(data))
    else:
        df = pd.read_excel(io.BytesIO(data))
    parsed = ParsedDocument(filename=filename, content_type=content_type, source_type="dataset")
    parsed.tables.append(
        {
            "columns": list(df.columns),
            "shape": list(df.shape),
            "records": df.head(200).astype(object).where(pd.notna(df.head(200)), None).to_dict("records"),
        }
    )
    summary = build_dataset_text_summary(filename, df)
    parsed.pages.append(ParsedPage(number=1, text=summary))
    return parsed


def build_dataset_text_summary(filename: str, df: pd.DataFrame) -> str:
    lines = [f"Dataset: {filename}", f"Rows: {len(df)}", f"Columns: {len(df.columns)}", ""]
    for column in df.columns:
        dtype = str(df[column].dtype)
        try:
            unique = int(df[column].nunique(dropna=True))
        except Exception:
            unique = -1
        sample = df[column].dropna().astype(str).head(3).tolist()
        lines.append(f"- {column} ({dtype}): {unique} unique values; sample: {', '.join(sample[:3])}")
    if len(df) > 0:
        lines.append("")
        lines.append("First rows:")
        lines.append(df.head(5).to_string(index=False))
    return "\n".join(lines)


def parse_document(filename: str, content_type: str, data: bytes) -> ParsedDocument:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return parse_pdf(data, filename, content_type)
    if lower.endswith(".docx"):
        return parse_docx(data, filename, content_type)
    if any(lower.endswith(ext) for ext in DATASET_EXTENSIONS):
        return parse_dataset(data, filename, content_type)
    if any(lower.endswith(ext) for ext in TEXT_EXTENSIONS):
        return parse_text(data, filename, content_type)
    if any(lower.endswith(ext) for ext in IMAGE_EXTENSIONS):
        return ParsedDocument(
            filename=filename,
            content_type=content_type,
            source_type="image",
            needs_vision=True,
        )
    if any(lower.endswith(ext) for ext in AUDIO_EXTENSIONS):
        return ParsedDocument(
            filename=filename,
            content_type=content_type,
            source_type="audio",
            needs_audio=True,
        )
    raise ValueError(f"unsupported file type: {filename}")
